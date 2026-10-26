"""Chamber temperature/humidity from a Tapo T310/T315.

Two routes:
1. parse_export(): the file the Tapo app e-mails (History → View All → download).
   TP-Link does not document its layout, so the parser looks for the columns by name
   (time, temperature, humidity; English or German), copes with CSV or Excel, comma or
   semicolon, decimal commas, a few lines of preamble, and °F.
2. tapo_sync.py: straight from the H100 hub on the home network (last 24 h only).

Both end in import_readings(), which stores local wall-clock times like manual readings.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from bisect import bisect_left
from datetime import date, datetime, time, timedelta, timezone

from . import db

TIME_WORDS = ("time", "date", "zeit", "datum", "tid", "fecha", "heure")
TEMP_WORDS = ("temp",)
HUM_WORDS = ("humid", "feucht", "rh", "luftf", "fukt")


@dataclass
class Parsed:
    rows: list[tuple[datetime, float | None, float | None]] = field(default_factory=list)
    skipped: int = 0                 # lines that had no usable time or no values
    fahrenheit: bool = False
    columns: tuple[str, str | None, str | None] = ("", None, None)

    @property
    def first(self):
        return self.rows[0][0] if self.rows else None

    @property
    def last(self):
        return self.rows[-1][0] if self.rows else None


def _norm(s) -> str:
    return str(s or "").strip().lower()


def _find(header: list[str], words) -> int | None:
    for i, h in enumerate(header):
        hn = _norm(h)
        if any(w in hn for w in words):
            return i
    return None


def _num(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\u00a0", "")
    for unit in ("°C", "°F", "\u2103", "\u2109", "%"):
        s = s.replace(unit, "")
    s = s.strip()
    if not s or s in ("-", "--", "N/A", "n/a"):
        return None
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


_DATE_RE = re.compile(r"^\s*(\d{1,4})[./-](\d{1,2})[./-](\d{1,4})")


def _dayfirst(samples: list[str]) -> bool:
    """Decide dd/mm vs mm/dd from the data itself; default day-first (Europe)."""
    for s in samples:
        m = _DATE_RE.match(s)
        if not m or len(m.group(1)) == 4:          # ISO yyyy-mm-dd: order is fixed
            continue
        a, b = int(m.group(1)), int(m.group(2))
        if a > 12:
            return True
        if b > 12:
            return False
    return True


def _parse_time(v, dayfirst: bool) -> datetime | None:
    if isinstance(v, datetime):
        return v.replace(tzinfo=None, microsecond=0)
    s = str(v or "").strip()
    if not s:
        return None
    s = s.replace("T", " ").rstrip("Z")
    fmts = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M"]
    dm = ["%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M",
          "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d/%m/%Y %I:%M %p", "%d/%m/%Y %I:%M:%S %p"]
    md = [f.replace("%d/%m", "%m/%d").replace("%d.%m", "%m.%d").replace("%d-%m", "%m-%d") for f in dm]
    for f in fmts + (dm + md if dayfirst else md + dm):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s).replace(tzinfo=None, microsecond=0)
    except ValueError:
        return None


def _table(data: bytes, filename: str) -> list[list]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):          # .xlsm: cells only, macros are never run
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if text.count(";") > text.count(",") else ","
    return [r for r in csv.reader(io.StringIO(text), delimiter=delim)]


def parse_export(data: bytes, filename: str) -> Parsed:
    """Raises ValueError with a readable message if no time + temperature/humidity columns are found."""
    rows = _table(data, filename)
    head_i = None
    for i, r in enumerate(rows[:30]):              # header may follow a few lines of preamble
        h = [str(c) for c in r]
        if _find(h, TEMP_WORDS) is not None or _find(h, HUM_WORDS) is not None:
            head_i = i
            break
    if head_i is None:
        raise ValueError("Couldn't find temperature or humidity columns in this file. "
                         "Is it the export from the Tapo app?")
    header = [str(c) for c in rows[head_i]]
    ci, hi = _find(header, TEMP_WORDS), _find(header, HUM_WORDS)
    ti = _find(header, TIME_WORDS)
    if ti is None and 0 not in (ci, hi):
        ti = 0          # Tapo export: the first column's heading is the date range, e.g. "2026/09/20 ... - ..."
    if ti is None:
        raise ValueError("Couldn't find the time column in this file. Is it the export from the Tapo app?")
    # a separate "Time" column next to "Date": combine them
    ti2 = None
    if ti is not None:
        rest = [i for i, h in enumerate(header) if i not in (ti, ci, hi) and any(w in _norm(h) for w in TIME_WORDS)]
        if rest and not any(w in _norm(header[ti]) for w in ("time", "zeit")):
            ti2 = rest[0]
    fahrenheit = ci is not None and any(x in _norm(header[ci]) for x in ("°f", "℉", "fahrenheit"))
    body = rows[head_i + 1:]

    def cell(r, i):
        return r[i] if i is not None and i < len(r) else None

    samples = [str(cell(r, ti) or "") for r in body[:200]]
    dayfirst = _dayfirst(samples)
    out = Parsed(fahrenheit=fahrenheit, columns=(header[ti], cell(header, ci), cell(header, hi)))
    for r in body:
        tv = cell(r, ti)
        if ti2 is not None:
            tv = f"{cell(r, ti)} {cell(r, ti2)}"
        t = _parse_time(tv, dayfirst)
        temp, rh = _num(cell(r, ci)), _num(cell(r, hi))
        if t is None or (temp is None and rh is None):
            if any(str(c or "").strip() for c in r):
                out.skipped += 1
            continue
        if fahrenheit and temp is not None:
            temp = (temp - 32) * 5 / 9
        out.rows.append((t, round(temp, 2) if temp is not None else None, rh))
    out.rows = merge_same_time(out.rows)
    if not out.rows:
        raise ValueError("The file has the right columns but no readings could be read from it.")
    return out


def merge_same_time(rows):
    """Readings are stored in local wall-clock time. When the clocks go back (last Sunday of October)
    02:00-02:59 happens twice, so two readings share one local time: average them instead of letting
    one silently replace the other (audit 26 Oct, #2). Also sorts by time."""
    by_t: dict[datetime, list] = {}
    for t, temp, rh in rows:
        by_t.setdefault(t, []).append((temp, rh))
    out = []
    for t in sorted(by_t):
        v = by_t[t]
        if len(v) == 1:
            out.append((t, *v[0]))
            continue
        temps = [x[0] for x in v if x[0] is not None]
        rhs = [x[1] for x in v if x[1] is not None]
        out.append((t, round(sum(temps) / len(temps), 2) if temps else None,
                    round(sum(rhs) / len(rhs), 1) if rhs else None))
    return out


def wall_hours(a: datetime, b: datetime) -> int:
    """Clock hours that exist between local times a and b. The hour skipped in March doesn't count,
    and the repeated October hour counts once (its readings are merged), so a full period is 100 %."""
    n, t = 0, a
    while t < b:
        aware = t.replace(tzinfo=db.TZ)
        if aware.astimezone(timezone.utc).astimezone(db.TZ).replace(tzinfo=None) == t:
            n += 1
        t += timedelta(hours=1)
    return n


def leave_out(rows, start: datetime | None, end: datetime | None):
    """Drop readings from start to end inclusive (e.g. the sensor was out of the chamber)."""
    if not start or not end:
        return rows, 0
    kept = [r for r in rows if not start <= r[0] <= end]
    return kept, len(rows) - len(kept)


def import_readings(con, chamber_id: int, rows, source: str) -> tuple[int, int]:
    """Store readings; returns (stored, rejected as out of range). Re-importing the same
    times replaces them, so importing overlapping exports never duplicates."""
    stored = rejected = 0
    for t, temp, rh in rows:
        try:
            db.add_chamber_reading(con, chamber_id, t.isoformat(timespec="seconds"), temp, rh, source=source)
            stored += 1
        except ValueError:
            rejected += 1
    return stored, rejected


def import_export(con, chamber_id: int, rows, lo: datetime | None = None, hi: datetime | None = None):
    """Import a Tapo export. With a leave-out period (lo..hi, inclusive), sensor readings already stored
    for that period are removed too: the sensor was out of the chamber, whichever way they arrived
    (audit 26 Oct, #9). Readings inside locked drying periods are kept. Returns (stored, rejected, removed)."""
    removed = 0
    if lo and hi:
        for src in db.SENSOR_SOURCES:
            removed += db.delete_chamber_readings(con, chamber_id, src, lo, hi)[0]
    stored, rejected = import_readings(con, chamber_id, leave_out(rows, lo, hi)[0], "tapo_export")
    return stored, rejected, removed


WEIGH_IN_TIME = time(12, 0)          # weigh-ins have a date only; shown at noon on the charts
FULL_DAY_HOURS = 12                  # a day average from fewer clock hours than this is marked partial


def weigh_in_time(day: date) -> datetime:
    return datetime.combine(day, WEIGH_IN_TIME)


def day_averages(days: list[date], chamber_rows) -> dict[date, tuple[float | None, float | None, int, bool] | None]:
    """For each weigh-in day: (average °C, average RH, number of readings, partial day) over all chamber
    readings on that calendar day, or None when the sensor recorded nothing that day.
    'partial' = readings cover fewer than FULL_DAY_HOURS different clock hours (sensor started late, gaps)."""
    by_day: dict[date, list] = {}
    for r in chamber_rows:
        t = datetime.fromisoformat(r["ts"])
        by_day.setdefault(t.date(), []).append((t, r["temp_c"], r["rh_pct"]))
    out = {}
    for d in days:
        v = by_day.get(d)
        if not v:
            out[d] = None
            continue
        temps = [x[1] for x in v if x[1] is not None]
        rhs = [x[2] for x in v if x[2] is not None]
        hours = len({x[0].hour for x in v})
        out[d] = (round(sum(temps) / len(temps), 1) if temps else None,
                  round(sum(rhs) / len(rhs), 0) if rhs else None,
                  len(v), hours < FULL_DAY_HOURS)
    return out


PERIOD_MIN_COVERAGE = 0.5            # below this share of hours with data, a period average is marked partial


def period_averages(days: list[date], chamber_rows) -> dict[date, tuple[float | None, float | None, int, bool, float] | None]:
    """For each weigh-in after the first: average °C and RH over the period since the previous weigh-in
    (12:00 previous day to 12:00 this day), with (readings, partial, coverage share of hours with data).
    None when the sensor recorded nothing in the period; the first weigh-in has no period and is left out."""
    pts = sorted((datetime.fromisoformat(r["ts"]), r["temp_c"], r["rh_pct"]) for r in chamber_rows)
    times = [p[0] for p in pts]
    days = sorted(days)
    out = {}
    for prev, cur in zip(days, days[1:]):
        a, b = weigh_in_time(prev), weigh_in_time(cur)
        v = pts[bisect_left(times, a):bisect_left(times, b)]        # from a (inclusive) to b (exclusive)
        if not v:
            out[cur] = None
            continue
        temps = [x[1] for x in v if x[1] is not None]
        rhs = [x[2] for x in v if x[2] is not None]
        hours_total = max(wall_hours(a, b), 1)
        covered = len({(x[0].date(), x[0].hour) for x in v}) / hours_total
        out[cur] = (round(sum(temps) / len(temps), 1) if temps else None,
                    round(sum(rhs) / len(rhs), 0) if rhs else None,
                    len(v), covered < PERIOD_MIN_COVERAGE, round(min(covered, 1.0), 2))
    return out


def aggregate(chamber_rows, hours: int) -> list[dict]:
    """Average chamber readings per block of `hours` (1, 12 or 24), blocks aligned to midnight.
    Each block: mid (time plotted), avg/min/max for temperature and humidity, n readings.
    Blocks without readings are left out, so charts can show the gap."""
    blocks: dict[datetime, list] = {}
    for r in chamber_rows:
        t = datetime.fromisoformat(r["ts"])
        start = datetime.combine(t.date(), time(t.hour - t.hour % hours))
        blocks.setdefault(start, []).append((r["temp_c"], r["rh_pct"]))
    out = []
    for start in sorted(blocks):
        v = blocks[start]
        temps = [x[0] for x in v if x[0] is not None]
        rhs = [x[1] for x in v if x[1] is not None]
        out.append({
            "start": start, "mid": start + timedelta(hours=hours) / 2, "n": len(v),
            "t": round(sum(temps) / len(temps), 2) if temps else None,
            "t_min": min(temps) if temps else None, "t_max": max(temps) if temps else None,
            "h": round(sum(rhs) / len(rhs), 1) if rhs else None,
            "h_min": min(rhs) if rhs else None, "h_max": max(rhs) if rhs else None,
        })
    return out
